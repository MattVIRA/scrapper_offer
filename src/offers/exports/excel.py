"""Excel file of the offers.

Rules:
- the file is never overwritten;
- new offers are added at the bottom;
- rows already present are neither moved nor rewritten, except the
  "En ligne", "Score" and "Mots trouvés" columns, and "Source" when it
  is empty;
- the user's columns are never modified.

Columns are found by their title, not by their position: the user can
sort the file, move or add columns.
"""

import os
import shutil
import zipfile
from dataclasses import dataclass
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from offers.errors import OffersError
from offers.sources import source_label

# Columns filled by the program.
# fmt: off
SCRIPT_COLUMNS = [
    "Réf.", "Titre", "Entreprise", "Pays", "Ville", "Durée", "Publiée le",
    "Ajoutée le", "En ligne", "Score", "Mots trouvés", "Lien", "Recherche", "Source",
]
# fmt: on
# Columns reserved for the user. The program only writes their title.
USER_COLUMNS = ["Statut", "Postulé le", "Notes"]

STATUS_CHOICES = ["À voir", "Intéressant", "Postulé", "Refusé", "Pas pour moi"]
ONLINE = "Oui"
NO_LONGER_ONLINE = "Plus en ligne"
LINK_TEXT = "Ouvrir l'offre"
SHEET_TITLE = "Offres"

# fmt: off
WIDTHS = {
    "Réf.": 12, "Titre": 55, "Entreprise": 28, "Pays": 22, "Ville": 18,
    "Durée": 8, "Publiée le": 12, "Ajoutée le": 12, "En ligne": 14,
    "Score": 7, "Mots trouvés": 28, "Lien": 16, "Recherche": 22, "Source": 16,
    "Statut": 14, "Postulé le": 12, "Notes": 40,
}
# fmt: on
DATE_FORMAT = "DD/MM/YYYY"
DATE_COLUMNS = ("Publiée le", "Ajoutée le")
# Light green background of the offers added today.
NEW_COLOR = "FFE2F0D9"
NEW_FILL = PatternFill("solid", fgColor=NEW_COLOR)
# The "Statut" drop-down list is set down to this row.
STATUS_LAST_ROW = 20000
BACKUP_DIR = "backup"
LOCK_PREFIX = "~$"


class ExcelError(OffersError):
    """Problem with the Excel file."""


@dataclass(frozen=True)
class ExcelReport:
    """What an update changed in the file."""

    added: int
    relevant: int
    removed: int


def check_file_closed(path):
    """Stop if the file is open in Excel.

    When Excel opens a file, it creates a "~$name.xlsx" file next to it.
    """
    folder, name = os.path.split(path)
    if os.path.exists(os.path.join(folder, LOCK_PREFIX + name)):
        raise ExcelError(f"{name} est ouvert : ferme l'Excel puis relance.")


def save_backup(path):
    """Copy the current file into the "backup" folder."""
    backup_dir = os.path.join(os.path.dirname(path), BACKUP_DIR)
    name, extension = os.path.splitext(os.path.basename(path))
    moment = datetime.now().strftime("%Y-%m-%d_%H%M")
    try:
        os.makedirs(backup_dir, exist_ok=True)
        shutil.copy2(path, os.path.join(backup_dir, f"{name}_{moment}{extension}"))
    except OSError as err:
        raise ExcelError(f"Impossible de sauvegarder {path} dans {backup_dir} : {err}") from err


def open_workbook(path):
    """Return the existing workbook, or a new one if the file does not exist."""
    if not os.path.exists(path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_TITLE
        sheet.freeze_panes = "A2"
        return workbook
    try:
        return load_workbook(path)
    except (OSError, zipfile.BadZipFile, InvalidFileException) as err:
        raise ExcelError(f"Impossible de lire {path} : {err}") from err


def save_workbook(workbook, path):
    """Save atomically: a temporary file first, then it replaces the old one.

    If saving fails, the old file stays intact.
    """
    temporary_path = path + ".tmp"
    try:
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    except OSError as err:
        if os.path.exists(temporary_path):
            os.remove(temporary_path)
        raise ExcelError(f"Impossible d'écrire {path} : {err}") from err


def read_columns(sheet):
    """Return the number of each column, from the titles of row 1."""
    columns = {}
    for cell in sheet[1]:
        if cell.value is not None:
            columns[str(cell.value).strip()] = cell.column
    return columns


def add_missing_columns(sheet, columns):
    """Add the missing columns on the right (all of them, if the file is new)."""
    for title in SCRIPT_COLUMNS + USER_COLUMNS:
        if title in columns:
            continue
        number = max(columns.values(), default=0) + 1
        cell = sheet.cell(row=1, column=number, value=title)
        cell.font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(number)].width = WIDTHS[title]
        columns[title] = number


def last_row(sheet):
    """Return the number of the last row that contains something."""
    row = sheet.max_row
    while row > 1 and all(c.value is None for c in sheet[row]):
        row -= 1
    return row


def write_new_row(sheet, columns, row, found, score, words, today, highlight):
    """Write a new offer on the given row."""
    offer = found.offer
    values = {
        "Réf.": offer.reference,
        "Titre": offer.title,
        "Entreprise": offer.company,
        "Pays": offer.country,
        "Ville": offer.city,
        "Durée": offer.duration_months,
        "Publiée le": offer.published_on,
        "Ajoutée le": today,
        "En ligne": ONLINE,
        "Score": score,
        "Mots trouvés": words,
        "Lien": LINK_TEXT,
        "Recherche": found.search_name,
        "Source": source_label(offer.source),
    }
    for title, value in values.items():
        cell = sheet.cell(row=row, column=columns[title], value=value)
        if title in DATE_COLUMNS:
            cell.number_format = DATE_FORMAT
        if title == "Lien":
            cell.hyperlink = offer.url
            cell.style = "Hyperlink"
        if highlight:
            cell.fill = NEW_FILL
    sheet.cell(row=row, column=columns["Postulé le"]).number_format = DATE_FORMAT


def remove_old_highlight(sheet, columns, end, today):
    """Remove the green background of the offers added before today.

    Only the program's green is removed; colors set by hand are kept.
    """
    for row in range(2, end + 1):
        added_on = sheet.cell(row=row, column=columns["Ajoutée le"]).value
        if isinstance(added_on, datetime):
            added_on = added_on.date()
        if added_on == today:
            continue
        for cell in sheet[row]:
            if cell.fill.fgColor.rgb == NEW_COLOR:
                cell.fill = PatternFill()


def update_existing_rows(sheet, columns, end, online_offers, can_remove, score):
    """Update "En ligne", "Score", "Mots trouvés" and an empty "Source".

    online_offers: offers known to be online, by reference.
    can_remove: True if online_offers is complete, so a row missing from it
        is no longer online.
    Return the references present and the number of rows marked as no
    longer online.
    """
    present = set()
    removed = 0
    for row in range(2, end + 1):
        ref = sheet.cell(row=row, column=columns["Réf."]).value
        if not ref:
            continue
        present.add(ref)
        online_cell = sheet.cell(row=row, column=columns["En ligne"])
        if ref in online_offers:
            offer = online_offers[ref]
            offer_score, words = score(offer)
            online_cell.value = ONLINE
            sheet.cell(row=row, column=columns["Score"]).value = offer_score
            sheet.cell(row=row, column=columns["Mots trouvés"]).value = words
            source_cell = sheet.cell(row=row, column=columns["Source"])
            if source_cell.value is None:
                source_cell.value = source_label(offer.source)
        elif can_remove and online_cell.value != NO_LONGER_ONLINE:
            online_cell.value = NO_LONGER_ONLINE
            removed += 1
    return present, removed


def add_status_list(sheet, columns):
    """Add the drop-down list of the "Statut" column."""
    letter = get_column_letter(columns["Statut"])
    choices = DataValidation(
        type="list", formula1='"' + ",".join(STATUS_CHOICES) + '"', allow_blank=True
    )
    choices.add(f"{letter}2:{letter}{STATUS_LAST_ROW}")
    sheet.add_data_validation(choices)


def update_workbook(path, found_offers, online_offers, score, today):
    """Create the Excel file, or add the new offers to the existing file.

    found_offers: FoundOffer list; only these offers can be added.
    online_offers: all offers online, by reference, to update the rows
        already present. None if unknown: then no row is marked as no
        longer online, and only rows found again are updated.
    score: function that returns (score, found words) for an Offer.
    today: date written in "Ajoutée le".
    """
    check_file_closed(path)
    file_exists = os.path.exists(path)
    workbook = open_workbook(path)
    sheet = workbook.active

    columns = read_columns(sheet)
    if file_exists and "Réf." not in columns:
        raise ExcelError(
            'La colonne "Réf." est introuvable dans le fichier Excel. '
            "Le script en a besoin pour reconnaître les offres."
        )
    add_missing_columns(sheet, columns)
    end = last_row(sheet)

    can_remove = online_offers is not None
    if not can_remove:
        online_offers = {found.offer.reference: found.offer for found in found_offers}
    present, removed = update_existing_rows(sheet, columns, end, online_offers, can_remove, score)
    remove_old_highlight(sheet, columns, end, today)

    # New offers, at the bottom. On the first run there is no green
    # background: everything would be green.
    added = 0
    relevant = 0
    for found in found_offers:
        ref = found.offer.reference
        if ref in present:
            continue
        present.add(ref)
        offer_score, words = score(found.offer)
        end += 1
        write_new_row(sheet, columns, end, found, offer_score, words, today, highlight=file_exists)
        added += 1
        if offer_score > 0:
            relevant += 1

    # The auto filter must also cover the new rows.
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{end}"
    if not file_exists:
        add_status_list(sheet, columns)

    # Backup last: nothing is written while an error is still possible.
    if file_exists:
        save_backup(path)
    save_workbook(workbook, path)
    return ExcelReport(added=added, relevant=relevant, removed=removed)
