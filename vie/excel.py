"""Création et mise à jour du fichier Excel des offres.

Règles :
- le fichier n'est jamais écrasé ;
- les nouvelles offres sont ajoutées en bas ;
- les lignes déjà présentes ne sont ni déplacées ni réécrites, sauf
  les colonnes "En ligne", "Score" et "Mots trouvés" ;
- les colonnes de l'utilisateur ne sont jamais modifiées.

Les colonnes sont retrouvées par leur titre, pas par leur position.
On peut donc trier le fichier, déplacer ou ajouter des colonnes.
"""

import os
import shutil
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from vie import api

# Colonnes remplies par le script.
COLONNES_SCRIPT = [
    "Réf.", "Titre", "Entreprise", "Pays", "Ville", "Durée", "Publiée le",
    "Ajoutée le", "En ligne", "Score", "Mots trouvés", "Lien", "Recherche",
]
# Colonnes réservées à l'utilisateur. Le script crée seulement leur titre.
COLONNES_UTILISATEUR = ["Statut", "Postulé le", "Notes"]

CHOIX_STATUT = ["À voir", "Intéressant", "Postulé", "Refusé", "Pas pour moi"]
EN_LIGNE = "Oui"
PLUS_EN_LIGNE = "Plus en ligne"

LARGEURS = {
    "Réf.": 12, "Titre": 55, "Entreprise": 28, "Pays": 22, "Ville": 18,
    "Durée": 8, "Publiée le": 12, "Ajoutée le": 12, "En ligne": 14,
    "Score": 7, "Mots trouvés": 28, "Lien": 16, "Recherche": 22,
    "Statut": 14, "Postulé le": 12, "Notes": 40,
}
FORMAT_DATE = "DD/MM/YYYY"
# Fond vert clair des offres ajoutées aujourd'hui.
COULEUR_NOUVEAU = "FFE2F0D9"
FOND_NOUVEAU = PatternFill("solid", fgColor=COULEUR_NOUVEAU)
# La liste déroulante "Statut" est posée jusqu'à cette ligne.
DERNIERE_LIGNE_STATUT = 20000


class ErreurExcel(Exception):
    """Problème avec le fichier Excel."""


def verifier_fichier_ferme(chemin):
    """Arrête le script si le fichier est ouvert dans Excel.

    Quand Excel ouvre un fichier, il crée à côté un fichier "~$nom.xlsx".
    """
    dossier, nom = os.path.split(chemin)
    if os.path.exists(os.path.join(dossier, "~$" + nom)):
        raise ErreurExcel(f"{nom} est ouvert : ferme l'Excel puis relance.")


def sauvegarder_copie(chemin):
    """Copie le fichier actuel dans le dossier "backup"."""
    dossier_copies = os.path.join(os.path.dirname(chemin), "backup")
    os.makedirs(dossier_copies, exist_ok=True)
    nom, extension = os.path.splitext(os.path.basename(chemin))
    moment = datetime.now().strftime("%Y-%m-%d_%H%M")
    shutil.copy2(chemin, os.path.join(dossier_copies, f"{nom}_{moment}{extension}"))


def lire_colonnes(feuille):
    """Renvoie le numéro de chaque colonne, à partir des titres de la ligne 1."""
    colonnes = {}
    for cellule in feuille[1]:
        if cellule.value is not None:
            colonnes[str(cellule.value).strip()] = cellule.column
    return colonnes


def completer_colonnes(feuille, colonnes):
    """Ajoute à droite les colonnes qui manquent (toutes, si le fichier est neuf)."""
    for titre in COLONNES_SCRIPT + COLONNES_UTILISATEUR:
        if titre in colonnes:
            continue
        numero = max(colonnes.values(), default=0) + 1
        cellule = feuille.cell(row=1, column=numero, value=titre)
        cellule.font = Font(bold=True)
        feuille.column_dimensions[get_column_letter(numero)].width = LARGEURS[titre]
        colonnes[titre] = numero


def derniere_ligne(feuille):
    """Renvoie le numéro de la dernière ligne qui contient quelque chose."""
    ligne = feuille.max_row
    while ligne > 1 and all(c.value is None for c in feuille[ligne]):
        ligne -= 1
    return ligne


def en_date(texte):
    """Transforme "2026-10-07T15:03:30Z" en date. Renvoie None si vide."""
    if not texte:
        return None
    return date.fromisoformat(texte[:10])


def en_nombre(texte):
    """Transforme "12" en 12. Garde le texte s'il n'est pas un nombre."""
    try:
        return int(texte)
    except (TypeError, ValueError):
        return texte


def ecrire_nouvelle_ligne(feuille, colonnes, ligne, offre, nom_recherche,
                          score, mots, aujourd_hui, colorer):
    """Écrit une nouvelle offre sur la ligne donnée."""
    valeurs = {
        "Réf.": api.reference(offre),
        "Titre": offre.get("missionTitle"),
        "Entreprise": offre.get("organizationName"),
        "Pays": offre.get("countryName"),
        "Ville": offre.get("cityName"),
        "Durée": en_nombre(offre.get("missionDuration")),
        "Publiée le": en_date(offre.get("startBroadcastDate") or offre.get("creationDate")),
        "Ajoutée le": aujourd_hui,
        "En ligne": EN_LIGNE,
        "Score": score,
        "Mots trouvés": mots,
        "Lien": "Ouvrir l'offre",
        "Recherche": nom_recherche,
    }
    for titre, valeur in valeurs.items():
        cellule = feuille.cell(row=ligne, column=colonnes[titre], value=valeur)
        if titre in ("Publiée le", "Ajoutée le"):
            cellule.number_format = FORMAT_DATE
        if titre == "Lien":
            cellule.hyperlink = api.lien(offre)
            cellule.style = "Hyperlink"
        if colorer:
            cellule.fill = FOND_NOUVEAU
    feuille.cell(row=ligne, column=colonnes["Postulé le"]).number_format = FORMAT_DATE


def enlever_ancien_fond(feuille, colonnes, fin, aujourd_hui):
    """Enlève le fond vert des offres ajoutées avant aujourd'hui.

    Seul le fond vert du script est enlevé. Les couleurs mises à la main
    ne sont pas touchées.
    """
    for ligne in range(2, fin + 1):
        ajout = feuille.cell(row=ligne, column=colonnes["Ajoutée le"]).value
        if isinstance(ajout, datetime):
            ajout = ajout.date()
        if ajout == aujourd_hui:
            continue
        for cellule in feuille[ligne]:
            if cellule.fill.fgColor.rgb == COULEUR_NOUVEAU:
                cellule.fill = PatternFill()


def mettre_a_jour(chemin, offres_site, offres_recherches, noter):
    """Crée le fichier Excel, ou ajoute les nouvelles offres au fichier existant.

    offres_site : toutes les offres du site, rangées par référence.
        Sert à savoir si une offre du fichier est encore en ligne.
    offres_recherches : liste de (offre, nom de la recherche).
        Seules ces offres peuvent être ajoutées au fichier.
    noter : fonction qui renvoie (score, mots trouvés) pour une offre.

    Renvoie un bilan : nombre d'offres ajoutées, pertinentes, retirées.
    """
    verifier_fichier_ferme(chemin)
    fichier_existe = os.path.exists(chemin)
    aujourd_hui = date.today()

    if fichier_existe:
        sauvegarder_copie(chemin)
        classeur = load_workbook(chemin)
        feuille = classeur.active
    else:
        os.makedirs(os.path.dirname(chemin), exist_ok=True)
        classeur = Workbook()
        feuille = classeur.active
        feuille.title = "Offres"
        feuille.freeze_panes = "A2"

    colonnes = lire_colonnes(feuille)
    if fichier_existe and "Réf." not in colonnes:
        raise ErreurExcel(
            'La colonne "Réf." est introuvable dans le fichier Excel. '
            "Le script en a besoin pour reconnaître les offres."
        )
    completer_colonnes(feuille, colonnes)
    fin = derniere_ligne(feuille)

    # 1) Lignes déjà présentes : on met à jour "En ligne", "Score" et
    #    "Mots trouvés". Rien d'autre.
    references_presentes = set()
    retirees = 0
    for ligne in range(2, fin + 1):
        ref = feuille.cell(row=ligne, column=colonnes["Réf."]).value
        if not ref:
            continue
        references_presentes.add(ref)
        cellule_en_ligne = feuille.cell(row=ligne, column=colonnes["En ligne"])
        if ref in offres_site:
            score, mots = noter(offres_site[ref])
            cellule_en_ligne.value = EN_LIGNE
            feuille.cell(row=ligne, column=colonnes["Score"]).value = score
            feuille.cell(row=ligne, column=colonnes["Mots trouvés"]).value = mots
        elif cellule_en_ligne.value != PLUS_EN_LIGNE:
            cellule_en_ligne.value = PLUS_EN_LIGNE
            retirees += 1

    enlever_ancien_fond(feuille, colonnes, fin, aujourd_hui)

    # 2) Nouvelles offres : ajoutées en bas du fichier.
    #    Au premier lancement, pas de fond vert : tout serait vert.
    ajoutees = 0
    pertinentes = 0
    for offre, nom_recherche in offres_recherches:
        ref = api.reference(offre)
        if ref in references_presentes:
            continue
        references_presentes.add(ref)
        score, mots = noter(offre)
        fin += 1
        ecrire_nouvelle_ligne(feuille, colonnes, fin, offre, nom_recherche,
                              score, mots, aujourd_hui, colorer=fichier_existe)
        ajoutees += 1
        if score > 0:
            pertinentes += 1

    # Le filtre automatique doit couvrir aussi les nouvelles lignes.
    feuille.auto_filter.ref = f"A1:{get_column_letter(feuille.max_column)}{fin}"

    if not fichier_existe:
        # Liste déroulante de la colonne "Statut".
        lettre = get_column_letter(colonnes["Statut"])
        liste = DataValidation(
            type="list", formula1='"' + ",".join(CHOIX_STATUT) + '"', allow_blank=True
        )
        liste.add(f"{lettre}2:{lettre}{DERNIERE_LIGNE_STATUT}")
        feuille.add_data_validation(liste)

    # On écrit d'abord un fichier temporaire, puis on remplace l'ancien.
    # Si l'écriture échoue, l'ancien fichier reste intact.
    chemin_temporaire = chemin + ".tmp"
    classeur.save(chemin_temporaire)
    os.replace(chemin_temporaire, chemin)

    return {"ajoutees": ajoutees, "pertinentes": pertinentes, "retirees": retirees}
