"""Tests of the Excel export (offers.exports.excel).

To regenerate the snapshot after a wanted change:
    UPDATE_SNAPSHOT=1 pytest tests/test_excel_export.py
"""

import json
import os
import re
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from offers.exports import excel
from offers.scoring.keywords import KeywordScorer
from offers.sources.base import FoundOffer
from offers.sources.business_france import to_offer

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
SNAPSHOT_PATH = os.path.join(FIXTURES, "excel_expected.json")
# Snapshot of the old code (T0.5), without the "Source" column.
OLD_SNAPSHOT_PATH = os.path.join(FIXTURES, "excel_attendu.json")
DAY_1 = date(2026, 10, 20)
DAY_2 = date(2026, 10, 21)
DAY_3 = date(2026, 10, 22)
SEARCH = "Recherche principale"
WORDS_DAY_1 = ["python", "data", "ia", "sql", "food"]
WORDS_DAY_2 = ["python", "data", "ia", "sql", "food", "marketing"]


# --- Helpers ------------------------------------------------------------------


def scorer(include, exclude=()):
    return KeywordScorer(include, exclude).score


def run(path, day, offers, words=WORDS_DAY_1, online=True):
    """Update the file as the pipeline does, with these offers online."""
    online_offers = {o.reference: o for o in offers} if online else None
    found = [FoundOffer(o, SEARCH) for o in offers]
    return excel.update_workbook(path, found, online_offers, scorer(words), day)


def rows_by_reference(path):
    """Return {reference: {column title: value}} for each row."""
    sheet = load_workbook(path).active
    titles = [c.value for c in sheet[1]]
    rows = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = dict(zip(titles, row))
        rows[values["Réf."]] = values
    return rows


def green_rows(path):
    """Return the references of the rows with the "new" green background."""
    sheet = load_workbook(path).active
    ref_column = [c.value for c in sheet[1]].index("Réf.")
    return [
        row[ref_column].value
        for row in sheet.iter_rows(min_row=2)
        if any(c.fill.fgColor.rgb == excel.NEW_COLOR for c in row)
    ]


def as_text(value):
    """Make a cell value storable in JSON."""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def content(path):
    """Return everything important in the file, for the snapshot.

    Keys stay those of the T0.5 snapshot, to compare with it.
    """
    sheet = load_workbook(path).active
    return {
        "titre_feuille": sheet.title,
        "volets_figes": sheet.freeze_panes,
        "filtre": sheet.auto_filter.ref,
        "listes": [
            {"formule": v.formula1, "plage": str(v.sqref)}
            for v in sheet.data_validations.dataValidation
        ],
        "cellules": [[as_text(c.value) for c in row] for row in sheet.iter_rows()],
        "liens": [
            [c.coordinate, c.hyperlink.target]
            for row in sheet.iter_rows()
            for c in row
            if c.hyperlink
        ],
        # Color written as is: the snapshot must see a color change.
        "fonds": [
            [c.coordinate, c.fill.fgColor.rgb]
            for row in sheet.iter_rows()
            for c in row
            if c.fill.fill_type == "solid"
        ],
    }


def drop_column(state, title):
    """Return the snapshot state as if the column did not exist."""
    number = state["cellules"][0].index(title) + 1

    def shift(coordinate, range_end=False):
        """Move a cell coordinate left if it is after the column.

        A cell in the column gives None, except at the end of a range,
        which then stops at the previous column.
        """
        letters, row = re.fullmatch(r"([A-Z]+)(\d+)", coordinate).groups()
        column = column_index_from_string(letters)
        if column == number and not range_end:
            return None
        return f"{get_column_letter(column - (column >= number))}{row}"

    def shift_range(cell_range):
        start, end = cell_range.split(":")
        return f"{shift(start)}:{shift(end, range_end=True)}"

    return {
        **state,
        "filtre": shift_range(state["filtre"]),
        "listes": [{**v, "plage": shift_range(v["plage"])} for v in state["listes"]],
        "cellules": [row[: number - 1] + row[number:] for row in state["cellules"]],
        "liens": [[shift(c), t] for c, t in state["liens"] if shift(c)],
        "fonds": [[shift(c), rgb] for c, rgb in state["fonds"] if shift(c)],
    }


def reorder_and_annotate(path):
    """Mimic a file edited by hand.

    - "Notes" column moved to the first position;
    - personal column "Ma colonne" added at the end;
    - Statut, Postulé le and Notes filled on two offers.
    Return the entered values, by reference.
    """
    old = load_workbook(path).active
    rows = [list(row) for row in old.iter_rows(values_only=True)]
    titles = rows[0]
    order = [titles.index("Notes")] + [i for i, t in enumerate(titles) if t != "Notes"]

    entered = {}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offres"
    for number, row in enumerate(rows):
        values = [row[i] for i in order]
        values.append("Ma colonne" if number == 0 else f"perso {number}")
        sheet.append(values)

    columns = excel.read_columns(sheet)
    for row, status, note in [(2, "Postulé", "Relancer lundi"), (3, "Intéressant", "Bof")]:
        ref = sheet.cell(row=row, column=columns["Réf."]).value
        sheet.cell(row=row, column=columns["Statut"]).value = status
        sheet.cell(row=row, column=columns["Postulé le"]).value = datetime(2026, 10, 20)
        sheet.cell(row=row, column=columns["Notes"]).value = note
        entered[ref] = {"Statut": status, "Notes": note}
    workbook.save(path)
    return entered


def two_days(path, offers):
    """Day 1 on a new file, day 2 on a file edited by hand."""
    result = {}
    run(path, DAY_1, offers[:4])
    result["jour_1"] = content(path)
    reorder_and_annotate(path)
    run(path, DAY_2, offers[1:], words=WORDS_DAY_2)
    result["jour_2"] = content(path)
    return result


@pytest.fixture
def offers(business_france_search):
    """The 5 offers of the fixture."""
    return [to_offer(raw) for raw in business_france_search["result"]]


@pytest.fixture
def path(tmp_path):
    return str(tmp_path / "offres_vie.xlsx")


# --- 1. Creation --------------------------------------------------------------


def test_file_creation(path, offers):
    report = run(path, DAY_1, offers)

    sheet = load_workbook(path).active
    assert [c.value for c in sheet[1]] == excel.SCRIPT_COLUMNS + excel.USER_COLUMNS
    assert sheet.title == "Offres"
    assert sheet.freeze_panes == "A2"

    rows = rows_by_reference(path)
    assert list(rows) == [o.reference for o in offers]
    first = rows[offers[0].reference]
    assert first["Titre"] == offers[0].title
    assert first["Entreprise"] == offers[0].company
    assert first["Ajoutée le"].date() == DAY_1
    assert first["En ligne"] == excel.ONLINE
    assert first["Recherche"] == SEARCH
    assert first["Source"] == "Business France"
    assert first["Statut"] is None and first["Notes"] is None

    # No green background on the first run: everything would be green.
    assert green_rows(path) == []

    lists = sheet.data_validations.dataValidation
    assert len(lists) == 1
    assert lists[0].formula1 == '"' + ",".join(excel.STATUS_CHOICES) + '"'

    assert sheet.cell(row=2, column=12).hyperlink.target == offers[0].url
    # No backup: there was nothing to save.
    assert not os.path.exists(os.path.join(os.path.dirname(path), "backup"))
    assert report.added == 5
    assert report.removed == 0
    assert report.relevant == sum(1 for o in offers if scorer(WORDS_DAY_1)(o)[0] > 0)


# --- 2. Update ----------------------------------------------------------------


def test_update_respects_the_existing_file(path, offers):
    run(path, DAY_1, offers[:4])
    entered = reorder_and_annotate(path)
    before = rows_by_reference(path)

    # Day 2: the 1st offer left the site, the 5th is new, and a keyword
    # was added.
    report = run(path, DAY_2, offers[1:], words=WORDS_DAY_2)
    after = rows_by_reference(path)

    sheet = load_workbook(path).active
    assert sheet.cell(row=1, column=1).value == "Notes"
    assert sheet.cell(row=1, column=sheet.max_column).value == "Ma colonne"

    # The user's entries and personal column are intact.
    for ref, values in entered.items():
        assert after[ref]["Statut"] == values["Statut"]
        assert after[ref]["Notes"] == values["Notes"]
        assert after[ref]["Postulé le"] == datetime(2026, 10, 20)
    for ref in before:
        assert after[ref]["Ma colonne"] == before[ref]["Ma colonne"]

    # Only "En ligne", "Score" and "Mots trouvés" change on existing rows.
    editable = {"En ligne", "Score", "Mots trouvés"}
    for ref in before:
        for title, value in before[ref].items():
            if title not in editable:
                assert after[ref][title] == value, (ref, title)

    gone = offers[0].reference
    assert after[gone]["En ligne"] == excel.NO_LONGER_ONLINE
    # The score of an offer that is no longer online is not recomputed.
    assert after[gone]["Score"] == before[gone]["Score"]

    for offer in offers[1:4]:
        score, words = scorer(WORDS_DAY_2)(offer)
        assert after[offer.reference]["Score"] == score
        assert after[offer.reference]["Mots trouvés"] == (words or None)

    # The new offer is at the bottom, on a green background.
    new = offers[4].reference
    assert list(after)[-1] == new
    assert after[new]["Ajoutée le"].date() == DAY_2
    assert green_rows(path) == [new]

    backups = os.listdir(os.path.join(os.path.dirname(path), "backup"))
    assert len(backups) == 1 and backups[0].startswith("offres_vie_")
    assert report.added == 1
    assert report.removed == 1


def test_offer_already_offline_is_not_counted_twice(path, offers):
    run(path, DAY_1, offers)
    assert run(path, DAY_2, offers[1:]).removed == 1
    assert run(path, DAY_3, offers[1:]).removed == 0


def test_offer_back_online(path, offers):
    run(path, DAY_1, offers)
    run(path, DAY_2, offers[1:])
    run(path, DAY_3, offers)
    assert rows_by_reference(path)[offers[0].reference]["En ligne"] == excel.ONLINE


def test_green_background_disappears_the_next_day(path, offers):
    run(path, DAY_1, offers[:4])
    run(path, DAY_2, offers)
    assert green_rows(path) == [offers[4].reference]

    run(path, DAY_3, offers)
    assert green_rows(path) == []


def test_no_duplicate_between_searches(path, offers):
    online = {o.reference: o for o in offers}
    # The same offer found by two searches is added once, with the name of
    # the first search.
    found = [FoundOffer(offers[0], "A"), FoundOffer(offers[0], "B"), FoundOffer(offers[1], "B")]
    report = excel.update_workbook(path, found, online, scorer(WORDS_DAY_1), DAY_1)

    rows = rows_by_reference(path)
    assert len(rows) == 2
    assert rows[offers[0].reference]["Recherche"] == "A"
    assert report.added == 2


def test_unknown_online_offers(path, offers):
    """online_offers=None: nothing is marked offline, unfound rows are kept."""
    run(path, DAY_1, offers[:3])
    before = rows_by_reference(path)

    report = run(path, DAY_2, offers[2:], words=WORDS_DAY_2, online=False)
    after = rows_by_reference(path)

    assert report.removed == 0 and report.added == 2
    for offer in offers[:2]:
        assert after[offer.reference] == before[offer.reference]
    # A row found again is online, so its score is updated.
    assert after[offers[2].reference]["Score"] == scorer(WORDS_DAY_2)(offers[2])[0]


def test_empty_source_is_filled_on_existing_rows(path, offers):
    """A file made before the "Source" column gets it on its online rows."""
    run(path, DAY_1, offers[:2])
    sheet_book = load_workbook(path)
    sheet = sheet_book.active
    sheet.delete_cols(excel.read_columns(sheet)["Source"])
    sheet_book.save(path)

    run(path, DAY_2, offers[1:])
    rows = rows_by_reference(path)

    sheet = load_workbook(path).active
    assert sheet.cell(row=1, column=sheet.max_column).value == "Source"
    assert rows[offers[0].reference]["Source"] is None
    assert all(rows[o.reference]["Source"] == "Business France" for o in offers[1:])


# --- 3. Errors ----------------------------------------------------------------


def test_error_if_the_file_is_open(path, offers):
    run(path, DAY_1, offers[:2])
    before = content(path)
    lock = os.path.join(os.path.dirname(path), "~$offres_vie.xlsx")
    open(lock, "w").close()

    with pytest.raises(excel.ExcelError, match="est ouvert"):
        run(path, DAY_2, offers)
    assert content(path) == before


def test_error_if_reference_column_is_missing(path, offers):
    workbook = Workbook()
    workbook.active.append(["Titre", "Notes"])
    workbook.active.append(["Une offre", "à garder"])
    workbook.save(path)

    with pytest.raises(excel.ExcelError, match="Réf."):
        run(path, DAY_1, offers)
    sheet = load_workbook(path).active
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [
        ["Titre", "Notes"],
        ["Une offre", "à garder"],
    ]
    assert not os.path.exists(os.path.join(os.path.dirname(path), "backup"))


def test_error_if_the_file_is_not_an_excel_file(path, offers):
    with open(path, "w") as file:
        file.write("pas un classeur")
    with pytest.raises(excel.ExcelError, match="Impossible de lire"):
        run(path, DAY_1, offers)


def test_error_on_save_keeps_the_old_file(monkeypatch, path, offers):
    run(path, DAY_1, offers[:2])
    before = content(path)

    def failing_replace(source, target):
        raise PermissionError("disque protégé")

    monkeypatch.setattr(excel.os, "replace", failing_replace)
    with pytest.raises(excel.ExcelError, match="Impossible d'écrire.*disque protégé"):
        run(path, DAY_2, offers)
    assert content(path) == before
    assert not os.path.exists(path + ".tmp")


# --- 4. Snapshots -------------------------------------------------------------


def test_snapshot(path, offers):
    """Compare the file with the recorded snapshot, after two runs."""
    result = two_days(path, offers)

    if os.environ.get("UPDATE_SNAPSHOT"):
        with open(SNAPSHOT_PATH, "w", encoding="utf-8") as file:
            json.dump(result, file, ensure_ascii=False, indent=1)
            file.write("\n")

    with open(SNAPSHOT_PATH, encoding="utf-8") as file:
        assert result == json.load(file)


def test_same_file_as_the_old_code_apart_from_source(path, offers):
    """Without the "Source" column, the file is the one of the T0.5 snapshot."""
    result = two_days(path, offers)
    with open(OLD_SNAPSHOT_PATH, encoding="utf-8") as file:
        old = json.load(file)
    assert {day: drop_column(state, "Source") for day, state in result.items()} == old
