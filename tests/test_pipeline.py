"""Tests of the collection pipeline, with fake sources."""

import os
from datetime import datetime

import pytest
from openpyxl import load_workbook

from offers.config import Config
from offers.errors import OffersError
from offers.exports import excel
from offers.pipeline import run_collect
from offers.storage.database import connect, last_runs, list_offers
from tests.factories import FakeSource, make_offer
from tests.test_excel_import import edit

DAY_1 = datetime(2026, 10, 20, 8, 0)
DAY_2 = datetime(2026, 10, 21, 8, 0)
CONFIG = Config(sources=[], keywords_include=["python", "data"], keywords_exclude=[])


def offer(number, source="fake"):
    return make_offer(
        source=source,
        source_id=str(number),
        reference=f"{source}-{number}",
        title=f"Data {number}",
        url=f"https://example.com/{number}",
    )


def rows(path):
    """Return {reference: "En ligne" value}."""
    sheet = load_workbook(path).active
    columns = excel.read_columns(sheet)
    return {
        row[columns["Réf."] - 1]: row[columns["En ligne"] - 1]
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }


def online_in_db(db):
    """Return {reference: online} of the stored offers."""
    conn = connect(db)
    try:
        return {s.offer.reference: s.online for s in list_offers(conn)}
    finally:
        conn.close()


def runs(db):
    conn = connect(db)
    try:
        return last_runs(conn, 100)
    finally:
        conn.close()


@pytest.fixture
def path(tmp_path):
    return str(tmp_path / "output" / "offres_vie.xlsx")


@pytest.fixture
def db(tmp_path):
    return str(tmp_path / "data" / "offers.db")


@pytest.fixture
def collect(db, path):
    """run_collect with the test config, database and Excel file."""

    def collect(now, *sources):
        return run_collect(CONFIG, db, path, now, sources=list(sources))

    return collect


def test_nominal_case(collect, db, path):
    offers = [offer(1), offer(2), offer(3)]
    report = collect(DAY_1, FakeSource(found=offers[:2], online=offers))

    assert online_in_db(db) == {"fake-1": True, "fake-2": True}
    assert rows(path) == {"fake-1": excel.ONLINE, "fake-2": excel.ONLINE}
    (source,) = report.sources
    assert source.name == "fake" and source.error is None
    assert (source.saved.found, source.saved.added, source.saved.relevant) == (2, 2, 2)
    assert (report.excel.added, report.excel.removed) == (2, 0)
    assert report.sync.updated == 0


def test_a_failing_source_does_not_stop_the_others(collect, db, path):
    collect(
        DAY_1,
        FakeSource(found=[offer(1, "good")], online=[offer(1, "good")], name="good"),
        FakeSource(found=[offer(1, "bad")], online=[offer(1, "bad")], name="bad"),
    )

    good = FakeSource(found=[offer(2, "good")], online=[offer(2, "good")], name="good")
    bad = FakeSource(error=OffersError("site en panne"), name="bad")
    report = collect(DAY_2, bad, good)

    # "good" gives its complete list: its offer 1 goes offline, even if
    # "bad" failed. The offers of "bad" do not change.
    assert rows(path) == {
        "good-1": excel.NO_LONGER_ONLINE,
        "bad-1": excel.ONLINE,
        "good-2": excel.ONLINE,
    }
    errors = {s.name: s.error for s in report.sources}
    assert errors == {"bad": "site en panne", "good": None}
    assert report.excel.added == 1 and report.excel.removed == 1


def test_a_source_that_cannot_tell_what_is_online(collect, path):
    collect(DAY_1, FakeSource(found=[offer(1)], online=[offer(1)]))
    report = collect(DAY_2, FakeSource(found=[offer(2)], online=None))
    assert rows(path) == {"fake-1": excel.ONLINE, "fake-2": excel.ONLINE}
    assert report.sources[0].saved.removed == 0


def test_all_sources_failing(collect, db, path):
    collect(DAY_1, FakeSource(found=[offer(1)], online=[offer(1)]))
    with open(path, "rb") as file:
        before = file.read()

    sources = [
        FakeSource(error=OffersError("panne A"), name="a"),
        FakeSource(error=OffersError("panne B"), name="b"),
    ]
    with pytest.raises(OffersError, match=r"Aucune source(.|\n)*- a : panne A\n- b : panne B"):
        collect(DAY_2, *sources)
    with open(path, "rb") as file:
        assert file.read() == before
    assert not os.path.exists(os.path.join(os.path.dirname(path), "backup"))
    # Only the run log is written.
    assert [(r.source, r.status, r.error) for r in runs(db)[:2]] == [
        ("b", "error", "panne B"),
        ("a", "error", "panne A"),
    ]
    assert online_in_db(db) == {"fake-1": True}


def test_no_enabled_source(db, path):
    with pytest.raises(OffersError, match="Aucune source activée"):
        run_collect(CONFIG, db, path, DAY_1)


def test_open_excel_stops_before_anything(collect, db, path):
    os.makedirs(os.path.dirname(path))
    open(os.path.join(os.path.dirname(path), "~$offres_vie.xlsx"), "w").close()
    source = FakeSource(found=[offer(1)], online=[offer(1)])

    with pytest.raises(excel.ExcelError, match="est ouvert"):
        collect(DAY_1, source)
    assert source.calls == 0
    assert not os.path.exists(db)


def test_two_runs_add_no_duplicate(collect, db, path):
    for day in (DAY_1, DAY_2):
        report = collect(day, FakeSource(found=[offer(1), offer(2)], online=[offer(1), offer(2)]))

    assert list(online_in_db(db)) == ["fake-1", "fake-2"]
    assert len(load_workbook(path).active["A"]) == 3
    assert report.sources[0].saved.added == 0 and report.excel.added == 0


def test_each_run_logs_every_source_and_saves_the_database(collect, db):
    sources = [
        FakeSource(found=[offer(1)], online=[offer(1)]),
        FakeSource(error=OffersError("panne"), name="broken"),
    ]
    clock_times = iter([datetime(2026, 10, 21, 8, m) for m in range(4)])
    run_collect(CONFIG, db, os.path.join(os.path.dirname(db), "x.xlsx"), DAY_1, sources,
                clock=lambda: next(clock_times))  # fmt: skip
    collect(DAY_2, *sources)

    log = runs(db)
    assert [(r.source, r.status) for r in log] == [
        ("broken", "error"),
        ("fake", "ok"),
        ("broken", "error"),
        ("fake", "ok"),
    ]
    assert (log[3].started_at, log[3].finished_at) == (
        datetime(2026, 10, 21, 8, 0),
        datetime(2026, 10, 21, 8, 1),
    )
    assert (log[3].found, log[3].added, log[3].removed) == (1, 1, 0)
    # No backup before the first run: the database did not exist.
    backups = os.listdir(os.path.join(os.path.dirname(db), "backup"))
    assert backups == ["offers_2026-10-21_0800.db"]


def test_tracking_entered_in_the_excel_is_kept(collect, db, path):
    collect(DAY_1, FakeSource(found=[offer(1), offer(2)], online=[offer(1), offer(2)]))
    edit(path, {("fake-1", "Statut"): "Postulé", ("fake-1", "Notes"): "Relancer"})

    report = collect(DAY_2, FakeSource(found=[offer(1), offer(3)], online=[offer(1), offer(3)]))

    conn = connect(db)
    stored = {s.offer.reference: s for s in list_offers(conn)}
    assert (stored["fake-1"].status, stored["fake-1"].notes) == ("applied", "Relancer")
    assert report.sync.updated == 1
    sheet_rows = {r[0]: r for r in load_workbook(path).active.iter_rows(values_only=True)}
    columns = excel.read_columns(load_workbook(path).active)
    assert sheet_rows["fake-1"][columns["Statut"] - 1] == "Postulé"


def test_sources_created_from_the_config(monkeypatch, db, path):
    from offers import sources as registry
    from offers.config import SourceConfig

    class ConfiguredSource(FakeSource):
        def __init__(self, config):
            super().__init__(config, found=[offer(1)], online=[offer(1)])

    monkeypatch.setitem(registry.SOURCES, "fake", ConfiguredSource)
    config = Config(
        sources=[SourceConfig("fake", True, []), SourceConfig("off", False, [])],
        keywords_include=[],
        keywords_exclude=[],
    )
    report = run_collect(config, db, path, DAY_1)
    assert [s.name for s in report.sources] == ["fake"]
