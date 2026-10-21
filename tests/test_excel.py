"""Tests du fichier Excel des offres (vie/excel.py).

Ces tests figent le comportement actuel. Ils servent de preuve au lot 1 :
le nouveau code doit produire le même fichier.

Pour régénérer l'instantané après un changement voulu :
    MAJ_INSTANTANE=1 pytest tests/test_excel.py
"""

import json
import os
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from vie import api, excel, filters

CHEMIN_INSTANTANE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "fixtures", "excel_attendu.json"
)
JOUR_1 = date(2026, 10, 20)
JOUR_2 = date(2026, 10, 21)
JOUR_3 = date(2026, 10, 22)
RECHERCHE = "Recherche principale"
MOTS_JOUR_1 = ["python", "data", "ia", "sql", "food"]
MOTS_JOUR_2 = ["python", "data", "ia", "sql", "food", "marketing"]


# --- Outils -------------------------------------------------------------------


def fixer_date(monkeypatch, jour):
    """Fait croire à excel.py que nous sommes le jour donné."""

    class DateFixe(date):
        @classmethod
        def today(cls):
            return jour

    monkeypatch.setattr(excel, "date", DateFixe)


def noteur(mots_inclure, mots_exclure=()):
    """Fonction de score, comme celle de run.py."""
    return lambda offre: filters.noter(offre, mots_inclure, list(mots_exclure))


def lancer(monkeypatch, chemin, jour, offres, mots=MOTS_JOUR_1):
    """Lance une mise à jour du fichier, comme run.py, avec ces offres en ligne."""
    fixer_date(monkeypatch, jour)
    offres_site = {api.reference(o): o for o in offres}
    offres_recherches = [(o, RECHERCHE) for o in offres]
    return excel.mettre_a_jour(chemin, offres_site, offres_recherches, noteur(mots))


def lignes_par_reference(chemin):
    """Renvoie {référence: {titre de colonne: valeur}} pour chaque ligne."""
    feuille = load_workbook(chemin).active
    titres = [c.value for c in feuille[1]]
    lignes = {}
    for ligne in feuille.iter_rows(min_row=2, values_only=True):
        valeurs = dict(zip(titres, ligne))
        lignes[valeurs["Réf."]] = valeurs
    return lignes


def lignes_vertes(chemin):
    """Renvoie les références des lignes sur fond vert « nouveau »."""
    feuille = load_workbook(chemin).active
    colonne_ref = [c.value for c in feuille[1]].index("Réf.")
    return [
        ligne[colonne_ref].value
        for ligne in feuille.iter_rows(min_row=2)
        if any(c.fill.fgColor.rgb == excel.COULEUR_NOUVEAU for c in ligne)
    ]


def en_texte(valeur):
    """Rend une valeur de cellule enregistrable en JSON."""
    if isinstance(valeur, (date, datetime)):
        return valeur.isoformat()
    return valeur


def contenu(chemin):
    """Renvoie tout ce que le fichier contient d'important, pour l'instantané."""
    feuille = load_workbook(chemin).active
    return {
        "titre_feuille": feuille.title,
        "volets_figes": feuille.freeze_panes,
        "filtre": feuille.auto_filter.ref,
        "listes": [
            {"formule": v.formula1, "plage": str(v.sqref)}
            for v in feuille.data_validations.dataValidation
        ],
        "cellules": [
            [en_texte(c.value) for c in ligne] for ligne in feuille.iter_rows()
        ],
        "liens": [
            [c.coordinate, c.hyperlink.target]
            for ligne in feuille.iter_rows()
            for c in ligne
            if c.hyperlink
        ],
        # Couleur écrite en dur : l'instantané doit voir un changement de couleur.
        "fonds": [
            [c.coordinate, c.fill.fgColor.rgb]
            for ligne in feuille.iter_rows()
            for c in ligne
            if c.fill.fill_type == "solid"
        ],
    }


def reordonner_et_annoter(chemin):
    """Imite un fichier retouché à la main.

    - colonne « Notes » déplacée en première position ;
    - colonne perso « Ma colonne » ajoutée à la fin ;
    - Statut, Postulé le et Notes remplis sur deux offres.
    Renvoie les valeurs saisies, par référence.
    """
    ancienne = load_workbook(chemin).active
    lignes = [list(ligne) for ligne in ancienne.iter_rows(values_only=True)]
    titres = lignes[0]
    ordre = [titres.index("Notes")] + [i for i, t in enumerate(titres) if t != "Notes"]

    saisies = {}
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Offres"
    for numero, ligne in enumerate(lignes):
        valeurs = [ligne[i] for i in ordre]
        valeurs.append("Ma colonne" if numero == 0 else f"perso {numero}")
        feuille.append(valeurs)

    colonnes = excel.lire_colonnes(feuille)
    for ligne, statut, note in [(2, "Postulé", "Relancer lundi"), (3, "Intéressant", "Bof")]:
        ref = feuille.cell(row=ligne, column=colonnes["Réf."]).value
        feuille.cell(row=ligne, column=colonnes["Statut"]).value = statut
        feuille.cell(row=ligne, column=colonnes["Postulé le"]).value = datetime(2026, 10, 20)
        feuille.cell(row=ligne, column=colonnes["Notes"]).value = note
        saisies[ref] = {"Statut": statut, "Notes": note}
    classeur.save(chemin)
    return saisies


@pytest.fixture
def offres(recherche_business_france):
    """Les 5 offres de la fixture."""
    return recherche_business_france["result"]


@pytest.fixture
def chemin(tmp_path):
    return str(tmp_path / "offres_vie.xlsx")


# --- 1. Création --------------------------------------------------------------


def test_creation_du_fichier(monkeypatch, chemin, offres):
    bilan = lancer(monkeypatch, chemin, JOUR_1, offres)

    feuille = load_workbook(chemin).active
    assert [c.value for c in feuille[1]] == excel.COLONNES_SCRIPT + excel.COLONNES_UTILISATEUR
    assert feuille.title == "Offres"
    assert feuille.freeze_panes == "A2"

    lignes = lignes_par_reference(chemin)
    assert list(lignes) == [api.reference(o) for o in offres]
    premiere = lignes[api.reference(offres[0])]
    assert premiere["Titre"] == offres[0]["missionTitle"]
    assert premiere["Entreprise"] == offres[0]["organizationName"]
    assert premiere["Ajoutée le"].date() == JOUR_1
    assert premiere["En ligne"] == excel.EN_LIGNE
    assert premiere["Recherche"] == RECHERCHE
    assert premiere["Statut"] is None and premiere["Notes"] is None

    # Pas de fond vert au premier lancement : tout serait vert.
    assert lignes_vertes(chemin) == []

    listes = feuille.data_validations.dataValidation
    assert len(listes) == 1
    assert listes[0].formula1 == '"' + ",".join(excel.CHOIX_STATUT) + '"'

    assert feuille.cell(row=2, column=12).hyperlink.target == api.lien(offres[0])
    # Pas de copie de sauvegarde : il n'y avait rien à sauvegarder.
    assert not os.path.exists(os.path.join(os.path.dirname(chemin), "backup"))
    assert bilan["ajoutees"] == 5
    assert bilan["retirees"] == 0
    assert bilan["pertinentes"] == sum(1 for o in offres if noteur(MOTS_JOUR_1)(o)[0] > 0)


# --- 2. Mise à jour -----------------------------------------------------------


def test_mise_a_jour_respecte_le_fichier_existant(monkeypatch, chemin, offres):
    lancer(monkeypatch, chemin, JOUR_1, offres[:4])
    saisies = reordonner_et_annoter(chemin)
    avant = lignes_par_reference(chemin)

    # Jour 2 : la 1re offre a disparu du site, la 5e est nouvelle,
    # et un mot-clé a été ajouté.
    bilan = lancer(monkeypatch, chemin, JOUR_2, offres[1:], mots=MOTS_JOUR_2)
    apres = lignes_par_reference(chemin)

    feuille = load_workbook(chemin).active
    assert feuille.cell(row=1, column=1).value == "Notes"
    assert feuille.cell(row=1, column=feuille.max_column).value == "Ma colonne"

    # Les saisies de l'utilisateur et la colonne perso sont intactes.
    for ref, valeurs in saisies.items():
        assert apres[ref]["Statut"] == valeurs["Statut"]
        assert apres[ref]["Notes"] == valeurs["Notes"]
        assert apres[ref]["Postulé le"] == datetime(2026, 10, 20)
    for ref in avant:
        assert apres[ref]["Ma colonne"] == avant[ref]["Ma colonne"]

    # Seuls « En ligne », « Score » et « Mots trouvés » changent sur les
    # lignes existantes.
    modifiables = {"En ligne", "Score", "Mots trouvés"}
    for ref in avant:
        for titre, valeur in avant[ref].items():
            if titre not in modifiables:
                assert apres[ref][titre] == valeur, (ref, titre)

    disparue = api.reference(offres[0])
    assert apres[disparue]["En ligne"] == excel.PLUS_EN_LIGNE
    # Le score d'une offre qui n'est plus en ligne n'est pas recalculé.
    assert apres[disparue]["Score"] == avant[disparue]["Score"]

    for offre in offres[1:4]:
        score, mots = noteur(MOTS_JOUR_2)(offre)
        assert apres[api.reference(offre)]["Score"] == score
        assert apres[api.reference(offre)]["Mots trouvés"] == (mots or None)

    # La nouvelle offre est en bas, sur fond vert.
    nouvelle = api.reference(offres[4])
    assert list(apres)[-1] == nouvelle
    assert apres[nouvelle]["Ajoutée le"].date() == JOUR_2
    assert lignes_vertes(chemin) == [nouvelle]

    copies = os.listdir(os.path.join(os.path.dirname(chemin), "backup"))
    assert len(copies) == 1 and copies[0].startswith("offres_vie_")
    assert bilan["ajoutees"] == 1
    assert bilan["retirees"] == 1


def test_offre_deja_plus_en_ligne_pas_comptee_deux_fois(monkeypatch, chemin, offres):
    lancer(monkeypatch, chemin, JOUR_1, offres)
    assert lancer(monkeypatch, chemin, JOUR_2, offres[1:])["retirees"] == 1
    assert lancer(monkeypatch, chemin, JOUR_3, offres[1:])["retirees"] == 0


def test_offre_revenue_en_ligne(monkeypatch, chemin, offres):
    lancer(monkeypatch, chemin, JOUR_1, offres)
    lancer(monkeypatch, chemin, JOUR_2, offres[1:])
    lancer(monkeypatch, chemin, JOUR_3, offres)
    assert lignes_par_reference(chemin)[api.reference(offres[0])]["En ligne"] == excel.EN_LIGNE


def test_le_fond_vert_disparait_le_lendemain(monkeypatch, chemin, offres):
    lancer(monkeypatch, chemin, JOUR_1, offres[:4])
    lancer(monkeypatch, chemin, JOUR_2, offres)
    assert lignes_vertes(chemin) == [api.reference(offres[4])]

    lancer(monkeypatch, chemin, JOUR_3, offres)
    assert lignes_vertes(chemin) == []


def test_pas_de_doublon_entre_recherches(monkeypatch, chemin, offres):
    fixer_date(monkeypatch, JOUR_1)
    offres_site = {api.reference(o): o for o in offres}
    # La même offre trouvée par deux recherches n'est ajoutée qu'une fois,
    # avec le nom de la première recherche.
    offres_recherches = [(offres[0], "A"), (offres[0], "B"), (offres[1], "B")]
    bilan = excel.mettre_a_jour(chemin, offres_site, offres_recherches, noteur(MOTS_JOUR_1))

    lignes = lignes_par_reference(chemin)
    assert len(lignes) == 2
    assert lignes[api.reference(offres[0])]["Recherche"] == "A"
    assert bilan["ajoutees"] == 2


# --- 3. Erreurs ---------------------------------------------------------------


def test_erreur_si_le_fichier_est_ouvert(monkeypatch, chemin, offres):
    lancer(monkeypatch, chemin, JOUR_1, offres[:2])
    avant = contenu(chemin)
    verrou = os.path.join(os.path.dirname(chemin), "~$offres_vie.xlsx")
    open(verrou, "w").close()

    with pytest.raises(excel.ErreurExcel, match="est ouvert"):
        lancer(monkeypatch, chemin, JOUR_2, offres)
    assert contenu(chemin) == avant


def test_erreur_si_colonne_reference_absente(monkeypatch, chemin, offres):
    classeur = Workbook()
    classeur.active.append(["Titre", "Notes"])
    classeur.active.append(["Une offre", "à garder"])
    classeur.save(chemin)

    with pytest.raises(excel.ErreurExcel, match="Réf."):
        lancer(monkeypatch, chemin, JOUR_1, offres)
    feuille = load_workbook(chemin).active
    assert [[c.value for c in ligne] for ligne in feuille.iter_rows()] == [
        ["Titre", "Notes"],
        ["Une offre", "à garder"],
    ]


# --- 4. Instantané ------------------------------------------------------------


def test_instantane(monkeypatch, chemin, offres):
    """Compare le fichier à l'instantané enregistré, après deux lancements.

    Jour 1 : fichier neuf (liens, liste déroulante, volets figés).
    Jour 2 : sur un fichier retouché à la main (colonnes déplacées, saisies).
    """
    resultat = {}
    lancer(monkeypatch, chemin, JOUR_1, offres[:4])
    resultat["jour_1"] = contenu(chemin)
    reordonner_et_annoter(chemin)
    lancer(monkeypatch, chemin, JOUR_2, offres[1:], mots=MOTS_JOUR_2)
    resultat["jour_2"] = contenu(chemin)

    if os.environ.get("MAJ_INSTANTANE"):
        with open(CHEMIN_INSTANTANE, "w", encoding="utf-8") as fichier:
            json.dump(resultat, fichier, ensure_ascii=False, indent=1)
            fichier.write("\n")

    with open(CHEMIN_INSTANTANE, encoding="utf-8") as fichier:
        attendu = json.load(fichier)
    assert resultat == attendu
