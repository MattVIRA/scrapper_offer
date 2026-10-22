"""Tests of the import of the Excel file into the database."""

from datetime import date, datetime, time

import pytest
from openpyxl import load_workbook

from offers.exports import excel
from offers.scoring.keywords import KeywordScorer
from offers.sources.base import FoundOffer
from offers.sources.business_france import to_offer
from offers.storage.database import connect, list_offers, save_source_result
from offers.storage.excel_import import import_workbook, sync_tracking
from tests.test_excel_export import reorder_and_annotate

DAY_1 = date(2026, 10, 20)
DAY_2 = date(2026, 10, 21)
NOW = datetime(2026, 10, 22, 9, 30)
SCORE = KeywordScorer(["data"], []).score


def offers_of(search):
    """Return the offers of a Business France API response."""
    return [to_offer(raw) for raw in search["result"]]


@pytest.fixture
def offers(business_france_search):
    return offers_of(business_france_search)


@pytest.fixture
def path(tmp_path, offers):
    """Excel file of the 5 offers, the 1st one no longer online, annotated."""
    path = str(tmp_path / "offres_vie.xlsx")
    write_excel(path, DAY_1, offers[:4])
    write_excel(path, DAY_2, offers[1:])
    edit(path, annotations(offers))
    return path


def annotations(offers):
    """Follow-up entered by the user on two offers."""
    first, second = offers[0].reference, offers[1].reference
    return {
        (first, "Statut"): "Postulé",
        (first, "Postulé le"): datetime(2026, 10, 20),
        (first, "Notes"): "Relancer lundi",
        (second, "Statut"): "Intéressant",
        (second, "Notes"): "Bof",
    }


@pytest.fixture
def conn(tmp_path):
    return connect(tmp_path / "offers.db")


def write_excel(path, day, offers):
    """Collect these offers on that day, in a database of its own, and export."""
    conn = connect(f"{path}.db")
    found = [FoundOffer(o, "Recherche principale") for o in offers]
    save_source_result(conn, "business_france", found, offers, SCORE, datetime.combine(day, time()))
    excel.update_workbook(path, list_offers(conn), day)
    conn.close()


def edit(path, changes):
    """Change cells: {(reference, column title): value}."""
    workbook = load_workbook(path)
    sheet = workbook.active
    columns = excel.read_columns(sheet)
    for row in range(2, sheet.max_row + 1):
        ref = sheet.cell(row=row, column=columns["Réf."]).value
        for (target, title), value in changes.items():
            if ref == target:
                sheet.cell(row=row, column=columns[title]).value = value
    workbook.save(path)


def excel_rows(path):
    """Return {reference: {column title: value}}."""
    sheet = load_workbook(path).active
    titles = [c.value for c in sheet[1]]
    rows = [dict(zip(titles, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    return {row["Réf."]: row for row in rows}


def stored(conn):
    return {s.offer.reference: s for s in list_offers(conn)}


def tracking_rows(conn):
    """Return {offer_id: (status, applied_on, notes, updated_at)}."""
    rows = conn.execute("SELECT * FROM tracking")
    return {
        r["offer_id"]: (r["status"], r["applied_on"], r["notes"], r["updated_at"]) for r in rows
    }


def test_all_offers_and_tracking_are_imported(conn, path, offers):
    rows = excel_rows(path)
    report = import_workbook(conn, path, NOW)

    result = stored(conn)
    assert sorted(result) == sorted(rows)
    for offer in offers:
        row, imported = rows[offer.reference], result[offer.reference]
        assert imported.offer.source == "business_france"
        assert imported.offer.source_id == offer.source_id
        assert imported.offer.title == offer.title
        assert imported.offer.company == offer.company
        assert imported.offer.duration_months == offer.duration_months
        assert imported.offer.published_on == offer.published_on
        assert imported.offer.url == offer.url
        assert imported.offer.raw == {}
        assert imported.first_seen == row["Ajoutée le"]
        assert imported.last_seen == NOW
        assert imported.online == (row["En ligne"] == excel.ONLINE)
        assert imported.search_name == "Recherche principale"
        # The keyword score is computed at the next collection.
        assert imported.score is None
    assert not result[offers[0].reference].online

    first, second = result[offers[0].reference], result[offers[1].reference]
    assert (first.status, first.applied_on, first.notes) == (
        "applied",
        "2026-10-20",
        "Relancer lundi",
    )
    assert (second.status, second.applied_on, second.notes) == ("interesting", None, "Bof")
    assert all(result[o.reference].status is None for o in offers[2:])
    assert report.offers_added == 5 and report.tracking_added == 2
    assert report.skipped == [] and report.unknown_statuses == {}


def test_file_edited_by_hand(conn, tmp_path, offers):
    """Moved columns, a personal column, and links lost: the id comes from "Réf."."""
    path = str(tmp_path / "offres_vie.xlsx")
    write_excel(path, DAY_1, offers)
    entered = reorder_and_annotate(path)

    report = import_workbook(conn, path, NOW)

    result = stored(conn)
    assert {s.offer.source_id for s in result.values()} == {o.source_id for o in offers}
    for ref, values in entered.items():
        assert result[ref].notes == values["Notes"]
        assert result[ref].applied_on == "2026-10-20"
    assert report.offers_added == 5 and report.tracking_added == 2


def test_old_file_without_source(conn, path):
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.delete_cols(excel.read_columns(sheet)["Source"])
    workbook.save(path)

    import_workbook(conn, path, NOW)
    assert {s.offer.source for s in list_offers(conn)} == {"business_france"}


def test_unknown_status_is_kept_in_the_notes(conn, path, offers):
    ref = offers[3].reference
    edit(path, {(ref, "Statut"): "Entretien", (ref, "Notes"): "Mardi 10h"})
    other = offers[4].reference
    edit(path, {(other, "Statut"): "Entretien"})

    report = import_workbook(conn, path, NOW)

    result = stored(conn)
    assert result[ref].status is None
    assert result[ref].notes == "Statut d'origine : Entretien\nMardi 10h"
    assert result[other].notes == "Statut d'origine : Entretien"
    assert report.unknown_statuses == {"Entretien": 2}


def test_rows_without_reference_or_id(conn, path, offers):
    ref = offers[2].reference
    edit(path, {(ref, "Réf."): None, (offers[3].reference, "Réf."): "MANUEL-1"})
    workbook = load_workbook(path)
    sheet = workbook.active
    columns = excel.read_columns(sheet)
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=columns["Réf."]).value == "MANUEL-1":
            sheet.cell(row=row, column=columns["Lien"]).hyperlink = None
    workbook.save(path)

    report = import_workbook(conn, path, NOW)

    assert ref not in stored(conn) and "MANUEL-1" not in stored(conn)
    assert report.offers_added == 3
    assert report.skipped == ["MANUEL-1"]


def test_second_import_changes_nothing(conn, path, offers):
    import_workbook(conn, path, NOW)
    before = list_offers(conn)
    tracking = [tuple(r) for r in conn.execute("SELECT * FROM tracking")]

    # Even if the user changed a status meanwhile: the import does not
    # synchronize, it only adds what is missing.
    edit(path, {(offers[1].reference, "Statut"): "Refusé"})
    report = import_workbook(conn, path, datetime(2026, 10, 23))

    assert list_offers(conn) == before
    assert [tuple(r) for r in conn.execute("SELECT * FROM tracking")] == tracking
    assert report.offers_added == 0 and report.tracking_added == 0


def test_excel_file_is_not_modified(conn, path):
    with open(path, "rb") as file:
        before = file.read()
    import_workbook(conn, path, NOW)
    with open(path, "rb") as file:
        assert file.read() == before


def test_file_without_reference_column(conn, tmp_path):
    path = tmp_path / "autre.xlsx"
    workbook = excel.Workbook()
    workbook.active.append(["Titre"])
    workbook.save(path)
    with pytest.raises(excel.ExcelError, match="Réf."):
        import_workbook(conn, path, NOW)


# --- Synchronization of the follow-up (sync_tracking) -------------------------

LATER = datetime(2026, 10, 23, 18, 0)


@pytest.fixture
def imported(conn, path):
    """Database filled by the import of the Excel file."""
    import_workbook(conn, path, NOW)
    return conn


def test_sync_changed_status_note_and_date(imported, path, offers):
    first, third = offers[0].reference, offers[2].reference
    edit(
        path,
        {
            (first, "Statut"): "Refusé",
            (first, "Postulé le"): datetime(2026, 10, 21),
            (first, "Notes"): "Pas retenu",
            (third, "Statut"): "À rédiger",
        },
    )
    report = sync_tracking(imported, path, LATER)

    result = stored(imported)
    assert (result[first].status, result[first].applied_on, result[first].notes) == (
        "rejected",
        "2026-10-21",
        "Pas retenu",
    )
    assert result[third].status == "to_write"
    assert report.updated == 2


def test_sync_cleared_note(imported, path, offers):
    second = offers[1].reference
    edit(path, {(second, "Notes"): None})
    sync_tracking(imported, path, LATER)
    assert stored(imported)[second].notes is None
    assert stored(imported)[second].status == "interesting"


def test_sync_unchanged_rows_are_not_rewritten(imported, path, offers):
    before = tracking_rows(imported)
    edit(path, {(offers[1].reference, "Notes"): "Finalement oui"})

    report = sync_tracking(imported, path, LATER)

    after = tracking_rows(imported)
    changed = {key for key in after if after[key] != before.get(key)}
    assert len(changed) == 1
    assert after[changed.pop()][3] == LATER.isoformat()
    assert report.updated == 1
    # A second run without change writes nothing.
    assert sync_tracking(imported, path, datetime(2026, 10, 24)).updated == 0
    assert tracking_rows(imported) == after


def test_sync_unknown_status(imported, path, offers):
    ref = offers[3].reference
    edit(path, {(ref, "Statut"): "Entretien"})

    report = sync_tracking(imported, path, LATER)
    assert stored(imported)[ref].notes == "Statut d'origine : Entretien"
    assert report.unknown_statuses == {"Entretien": 1}
    assert sync_tracking(imported, path, LATER).updated == 0


def test_sync_missing_file(imported, tmp_path):
    report = sync_tracking(imported, str(tmp_path / "absent.xlsx"), LATER)
    assert report.updated == 0 and report.unknown_references == 0


def test_sync_unknown_reference(imported, path, offers):
    edit(path, {(offers[4].reference, "Réf."): "VIE1", (offers[4].reference, "Statut"): None})
    edit(path, {("VIE1", "Statut"): "Postulé"})

    report = sync_tracking(imported, path, LATER)
    assert report.unknown_references == 1 and report.updated == 0


def test_sync_with_moved_columns(conn, tmp_path, offers):
    path = str(tmp_path / "offres_vie.xlsx")
    write_excel(path, DAY_1, offers)
    import_workbook(conn, path, NOW)
    entered = reorder_and_annotate(path)

    report = sync_tracking(conn, path, LATER)

    result = stored(conn)
    for ref, values in entered.items():
        assert result[ref].notes == values["Notes"]
    assert report.updated == 2
